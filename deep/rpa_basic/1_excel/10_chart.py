# 학습: https://openpyxl.readthedocs.io/en/stable/charts/introduction.html

from openpyxl import load_workbook
wb = load_workbook("/Users/gunju/Desktop/self study/python/deep/rpa_basic/1_excel/sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart

# bar_value = Reference(ws, min_row = 2, max_row = 11, min_col = 2, max_col = 3) # 차트에 넣을 데이터 지정
# bar_chart = BarChart() # 차트 종류 설정 Bar, Line, Pie, ...
# bar_chart.add_data(bar_value) # 차트에 데이터 넣기
# ws.add_chart(bar_chart, "E1") # 차트 넣을 위치 지정

line_value = Reference(ws, min_row = 1, max_row = 11, min_col = 2, max_col = 3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data = True) # 계열 제목 지정
line_chart.title = "성적표" # 차트 이름
line_chart.style = 10 # 미리 정의된 스타일 적용, 사용자 개별 지정도 가능
line_chart.y_axis.title = "점수" # y축 제목
line_chart.x_axis.title = "번호" # x축 제목
ws.add_chart(line_chart, "E1")

wb.save("/Users/gunju/Desktop/self study/python/deep/rpa_basic/1_excel/sample_chart.xlsx")